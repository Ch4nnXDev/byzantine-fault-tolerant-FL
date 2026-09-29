import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


class ResearchReportGenerator:

    def __init__(
        self,
        metrics,
        plot_files
    ):

        self.metrics = metrics
        self.plot_files = plot_files
        self.output_dir = metrics.output_dir

    # =========================================================
    # Generate report
    # =========================================================

    def generate(self):

        document = Document()

        self.add_title(document)

        self.add_experiment_information(
            document
        )

        self.add_summary(
            document
        )

        self.add_convergence_analysis(
            document
        )

        self.add_runtime_analysis(
            document
        )

        self.add_figures(
            document
        )

        self.add_client_analysis(
            document
        )

        self.add_conclusion(
            document
        )

        path = os.path.join(
            self.output_dir,
            "research_report.docx"
        )

        document.save(path)

        return path

    # =========================================================
    # Title
    # =========================================================

    def add_title(
        self,
        document
    ):

        title = document.add_heading(
            "Federated Learning Experiment Report",
            level=0
        )

        title.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.add_run(
            self.metrics.experiment_name
        ).bold = True

        document.add_paragraph()

    # =========================================================
    # Experiment configuration
    # =========================================================

    def add_experiment_information(
        self,
        document
    ):

        document.add_heading(
            "1. Experiment Configuration",
            level=1
        )

        metadata = self.metrics.metadata

        table = document.add_table(
            rows=0,
            cols=2
        )

        table.style = (
            "Light Grid Accent 1"
        )

        fields = [

            (
                "Experiment",
                metadata.get(
                    "experiment_name"
                )
            ),

            (
                "Experiment ID",
                metadata.get(
                    "experiment_id"
                )
            ),

            (
                "Dataset",
                metadata.get(
                    "dataset"
                )
            ),

            (
                "Model",
                metadata.get(
                    "model"
                )
            ),

            (
                "Number of Clients",
                metadata.get(
                    "num_clients"
                )
            ),

            (
                "Number of Rounds",
                metadata.get(
                    "num_rounds"
                )
            ),

            (
                "Local Epochs",
                metadata.get(
                    "local_epochs"
                )
            ),

            (
                "Batch Size",
                metadata.get(
                    "batch_size"
                )
            ),

            (
                "Learning Rate",
                metadata.get(
                    "learning_rate"
                )
            ),

            (
                "Partition Type",
                metadata.get(
                    "partition_type"
                )
            ),

            (
                "Byzantine Clients",
                metadata.get(
                    "byzantine_clients"
                )
            ),

            (
                "Attack",
                metadata.get(
                    "attack"
                )
            ),

            (
                "Aggregation Method",
                metadata.get(
                    "aggregation_method"
                )
            ),

            (
                "Topology",
                metadata.get(
                    "topology"
                )
            ),

            (
                "Defense Enabled",
                metadata.get(
                    "defense_enabled"
                )
            ),

            (
                "Defense Method",
                metadata.get(
                    "defense_method"
                )
            ),

            (
                "Random Seed",
                metadata.get(
                    "random_seed"
                )
            )
        ]

        for name, value in fields:

            row = table.add_row().cells

            row[0].text = str(
                name
            )

            row[1].text = str(
                value
            )

    # =========================================================
    # Experiment summary
    # =========================================================

    def add_summary(
        self,
        document
    ):

        document.add_heading(
            "2. Experiment Summary",
            level=1
        )

        convergence = (
            self.metrics
            .get_convergence_metrics()
        )

        if not convergence:

            document.add_paragraph(
                "No experiment metrics "
                "were recorded."
            )

            return

        paragraph = (
            document.add_paragraph()
        )

        paragraph.add_run(
            "Initial Accuracy: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['initial_accuracy']:.2%}"
        )

        paragraph.add_run(
            "\nFinal Accuracy: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['final_accuracy']:.2%}"
        )

        paragraph.add_run(
            "\nBest Accuracy: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['best_accuracy']:.2%}"
        )

        paragraph.add_run(
            "\nBest Accuracy Round: "
        ).bold = True

        paragraph.add_run(
            str(
                convergence[
                    "best_accuracy_round"
                ]
            )
        )

        paragraph.add_run(
            "\nAccuracy Gain: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['accuracy_gain']:+.2%}"
        )

        paragraph.add_run(
            "\nInitial Loss: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['initial_loss']:.4f}"
        )

        paragraph.add_run(
            "\nFinal Loss: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['final_loss']:.4f}"
        )

        paragraph.add_run(
            "\nLoss Reduction: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['loss_reduction']:.4f}"
        )

        paragraph.add_run(
            "\nLoss Reduction Percentage: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['loss_reduction_percent']:.2f}%"
        )

    # =========================================================
    # Convergence analysis
    # =========================================================

    def add_convergence_analysis(
        self,
        document
    ):

        document.add_heading(
            "3. Convergence Analysis",
            level=1
        )

        convergence = (
            self.metrics
            .get_convergence_metrics()
        )

        if not convergence:
            return

        paragraph = (
            document.add_paragraph()
        )

        paragraph.add_run(
            "Largest round-to-round "
            "accuracy improvement: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['max_accuracy_improvement']:+.2%}"
        )

        paragraph.add_run(
            "\nAverage absolute "
            "round-to-round accuracy change: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['average_accuracy_change']:.4f}"
        )

        paragraph.add_run(
            "\nAccuracy volatility: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['accuracy_volatility']:.4f}"
        )

        paragraph.add_run(
            "\nAverage round-to-round "
            "loss change: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['average_loss_change']:.4f}"
        )

        paragraph.add_run(
            "\nLoss volatility: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['loss_volatility']:.4f}"
        )

        document.add_paragraph()

        document.add_paragraph(
            "These measurements describe the "
            "convergence behaviour and stability "
            "of the federated learning process "
            "across communication rounds."
        )

    # =========================================================
    # Runtime analysis
    # =========================================================

    def add_runtime_analysis(
        self,
        document
    ):

        document.add_heading(
            "4. Runtime Analysis",
            level=1
        )

        convergence = (
            self.metrics
            .get_convergence_metrics()
        )

        if not convergence:

            document.add_paragraph(
                "No runtime metrics "
                "were recorded."
            )

            return

        paragraph = (
            document.add_paragraph()
        )

        paragraph.add_run(
            "Average Round Duration: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['average_round_duration']:.2f} seconds"
        )

        paragraph.add_run(
            "\nMinimum Round Duration: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['min_round_duration']:.2f} seconds"
        )

        paragraph.add_run(
            "\nMaximum Round Duration: "
        ).bold = True

        paragraph.add_run(
            f"{convergence['max_round_duration']:.2f} seconds"
        )

        total_duration = (
            convergence[
                "total_experiment_duration"
            ]
        )

        if total_duration is not None:

            paragraph.add_run(
                "\nTotal Experiment Duration: "
            ).bold = True

            paragraph.add_run(
                f"{total_duration:.2f} seconds"
            )

        document.add_paragraph()

        document.add_paragraph(
            "Runtime measurements provide a "
            "baseline for evaluating computational "
            "overhead when additional Byzantine "
            "resilience, topology management, or "
            "defense mechanisms are introduced."
        )

    # =========================================================
    # Figures
    # =========================================================

    def add_figures(
        self,
        document
    ):

        document.add_heading(
            "5. Experimental Results",
            level=1
        )

        figures = [

            (
                "Global Accuracy",
                self.plot_files.get(
                    "accuracy"
                )
            ),

            (
                "Global Loss",
                self.plot_files.get(
                    "global_loss"
                )
            ),

            (
                "Round-to-Round Accuracy Change",
                self.plot_files.get(
                    "accuracy_change"
                )
            ),

            (
                "Client Training Loss",
                self.plot_files.get(
                    "client_losses"
                )
            ),

            (
                "Mean Client Loss",
                self.plot_files.get(
                    "client_mean_loss"
                )
            ),

            (
                "Client Loss Variability",
                self.plot_files.get(
                    "client_loss_variability"
                )
            ),

            (
                "Federated Round Duration",
                self.plot_files.get(
                    "round_duration"
                )
            )
        ]

        for title, path in figures:

            if not path:
                continue

            if not os.path.exists(path):
                continue

            document.add_heading(
                title,
                level=2
            )

            document.add_picture(
                path,
                width=Inches(6.2)
            )

            paragraph = (
                document.add_paragraph()
            )

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

    # =========================================================
    # Client-level analysis
    # =========================================================

    def add_client_analysis(
        self,
        document
    ):

        document.add_heading(
            "6. Client-Level Analysis",
            level=1
        )

        client_metrics = (
            self.metrics
            .get_client_metrics()
        )

        if not client_metrics:

            document.add_paragraph(
                "No client-level metrics "
                "were recorded."
            )

            return

        client_convergence = (
            self.metrics
            .get_client_convergence_metrics()
        )

        table = document.add_table(
            rows=1,
            cols=6
        )

        table.style = (
            "Light Grid Accent 1"
        )

        header = (
            table.rows[0].cells
        )

        header[0].text = "Client"
        header[1].text = "Rounds"
        header[2].text = "Initial Loss"
        header[3].text = "Final Loss"
        header[4].text = "Loss Change"
        header[5].text = "Loss Reduction %"

        for client_id, losses in (
            client_metrics.items()
        ):

            if not losses:
                continue

            data = client_convergence.get(
                client_id,
                {}
            )

            row = (
                table.add_row().cells
            )

            row[0].text = str(
                client_id
            )

            row[1].text = str(
                len(losses)
            )

            row[2].text = (
                f"{losses[0]:.4f}"
            )

            row[3].text = (
                f"{losses[-1]:.4f}"
            )

            row[4].text = (
                f"{data.get('loss_change', 0.0):+.4f}"
            )

            row[5].text = (
                f"{data.get('loss_reduction_percent', 0.0):.2f}%"
            )

    # =========================================================
    # Conclusion
    # =========================================================

    def add_conclusion(
        self,
        document
    ):

        document.add_heading(
            "7. Conclusion",
            level=1
        )

        convergence = (
            self.metrics
            .get_convergence_metrics()
        )

        if not convergence:

            document.add_paragraph(
                "The experiment did not produce "
                "sufficient metrics for automated "
                "analysis."
            )

            return

        initial_accuracy = (
            convergence[
                "initial_accuracy"
            ]
        )

        final_accuracy = (
            convergence[
                "final_accuracy"
            ]
        )

        accuracy_gain = (
            convergence[
                "accuracy_gain"
            ]
        )

        loss_reduction = (
            convergence[
                "loss_reduction"
            ]
        )

        # -----------------------------------------------------
        # Accuracy statement
        # -----------------------------------------------------

        if accuracy_gain > 0:

            accuracy_statement = (

                f"Global accuracy changed from "
                f"{initial_accuracy:.2%} to "
                f"{final_accuracy:.2%}, with a "
                f"net change of "
                f"{accuracy_gain:+.2%}."
            )

        elif accuracy_gain < 0:

            accuracy_statement = (

                f"Global accuracy changed from "
                f"{initial_accuracy:.2%} to "
                f"{final_accuracy:.2%}, with a "
                f"net change of "
                f"{accuracy_gain:+.2%}."
            )

        else:

            accuracy_statement = (

                f"Global accuracy remained at "
                f"{final_accuracy:.2%} between "
                f"the first and final recorded "
                f"rounds."
            )

        # -----------------------------------------------------
        # Loss statement
        # -----------------------------------------------------

        if loss_reduction > 0:

            loss_statement = (

                f"Global loss changed from "
                f"{convergence['initial_loss']:.4f} "
                f"to {convergence['final_loss']:.4f}, "
                f"representing a reduction of "
                f"{loss_reduction:.4f}."
            )

        elif loss_reduction < 0:

            loss_statement = (

                f"Global loss changed from "
                f"{convergence['initial_loss']:.4f} "
                f"to {convergence['final_loss']:.4f}, "
                f"representing an increase of "
                f"{abs(loss_reduction):.4f}."
            )

        else:

            loss_statement = (
                "Global loss remained unchanged "
                "between the first and final "
                "recorded rounds."
            )

        document.add_paragraph(
            accuracy_statement
        )

        document.add_paragraph(
            loss_statement
        )

        document.add_paragraph(
            "The report is generated directly "
            "from the recorded experiment "
            "configuration and metrics."
        )